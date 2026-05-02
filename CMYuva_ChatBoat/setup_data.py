#!/usr/bin/env python3
"""
setup_data.py - Run this ONCE to convert Excel files to JSON for fast loading.
Usage: python setup_data.py
Place all three .xlsx files in the same directory as this script.
"""
import json
import os
import sys
from openpyxl import load_workbook

DATA_DIR = os.path.join(os.path.dirname(__file__), 'data')
os.makedirs(DATA_DIR, exist_ok=True)

def process_excel(path, name, max_rows=None):
    print(f"Processing {name}...")
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    headers = None
    data = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            headers = [str(h) if h is not None else f'col_{i}' for i, h in enumerate(row)]
            continue
        row_dict = {}
        for j, h in enumerate(headers):
            v = row[j] if j < len(row) else None
            if hasattr(v, 'strftime'):
                v = v.strftime('%Y-%m-%d %H:%M:%S')
            elif isinstance(v, float) and v != v:  # NaN
                v = None
            row_dict[h] = v
        data.append(row_dict)
        if max_rows and len(data) >= max_rows:
            print(f"  (limited to {max_rows} rows for performance)")
            break
    wb.close()
    print(f"  ✓ {len(data)} rows loaded")
    return headers, data

def main():
    base = os.path.dirname(os.path.abspath(__file__))

    # Find xlsx files (flexible naming)
    def find_file(keywords):
        for f in os.listdir(base):
            if f.endswith('.xlsx') and all(k.lower() in f.lower() for k in keywords):
                return os.path.join(base, f)
        return None

    loan_path = find_file(['loan']) or find_file(['cmyuva', 'loan'])
    cert_path = find_file(['certif'])
    fraud_path = find_file(['fraud'])

    missing = []
    if not loan_path: missing.append('cmyuva-loan-applied-data.xlsx')
    if not cert_path: missing.append('cmyuva-certification-data.xlsx')
    if not fraud_path: missing.append('cm-yuva-fraud-data.xlsx')

    if missing:
        print("ERROR: Missing files:")
        for m in missing: print(f"  - {m}")
        print(f"\nPlace the .xlsx files in: {base}")
        sys.exit(1)

    print(f"\nFound files:")
    print(f"  Loan: {loan_path}")
    print(f"  Cert: {cert_path}")
    print(f"  Fraud: {fraud_path}")
    print()

    # Process loan (large file - take up to 30k rows)
    loan_headers, loan_data = process_excel(loan_path, 'Loan Data', max_rows=30000)
    with open(os.path.join(DATA_DIR, 'loan_data.json'), 'w', encoding='utf-8') as f:
        json.dump({'headers': loan_headers, 'data': loan_data}, f, ensure_ascii=False)

    # Process cert
    cert_headers, cert_data = process_excel(cert_path, 'Certification Data')
    with open(os.path.join(DATA_DIR, 'cert_data.json'), 'w', encoding='utf-8') as f:
        json.dump(cert_data, f, ensure_ascii=False)

    # Process fraud
    fraud_headers, fraud_data = process_excel(fraud_path, 'Fraud Data')
    # Clean keys
    cleaned_fraud = []
    for r in fraud_data:
        cleaned = {}
        for k, v in r.items():
            clean_k = k.strip().replace('\n', ' ').replace('  ', ' ')
            cleaned[clean_k] = str(v).strip() if v is not None else None
        cleaned_fraud.append(cleaned)
    with open(os.path.join(DATA_DIR, 'fraud_data.json'), 'w', encoding='utf-8') as f:
        json.dump(cleaned_fraud, f, ensure_ascii=False)

    print("\n✅ All data processed successfully!")
    print(f"   Saved to: {DATA_DIR}")
    print("\nNow run: python backend/app.py")

if __name__ == '__main__':
    main()
